import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config.settings import REPORTS_DIR

_PURPLE    = "5B3FA0"
_LIGHT_PUR = "D4C5EE"
_GREEN     = "C6EFCE"
_RED       = "FFC7CE"
_YELLOW    = "FFEB9C"
_WHITE     = "FFFFFF"
_GRAY      = "F2F2F2"
_BD = Border(
    left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),  bottom=Side(style="thin", color="CCCCCC"),
)


def _cell(ws, row, col, value, bg=_WHITE, bold=False, color="000000", center=True):
    c = ws.cell(row, col, value)
    c.font      = Font(name="Arial", bold=bold, color=color)
    c.fill      = PatternFill("solid", start_color=bg)
    c.alignment = Alignment(horizontal="center" if center else "left", vertical="center")
    c.border    = _BD
    return c


def analyze_cumulative_flow(records: list) -> dict:
    valid = [r for r in records if r.pass_fail in ("PASS", "FAIL")]
    if not valid:
        return {}
    initial      = valid[0].balance_before
    final        = valid[-1].balance_after
    total_bet    = sum(r.bet for r in valid)
    total_win    = sum(r.win_displayed for r in valid)
    expected_fin = initial - total_bet + total_win
    ok           = final == expected_fin
    return {
        "初始資產":    initial,
        "最終資產":    final,
        "總押注":      total_bet,
        "總贏分":      total_win,
        "預期最終資產": expected_fin,
        "累積金流正確": "✓ PASS" if ok else f"✗ FAIL（差 {final - expected_fin:+d}）",
    }


def export_excel(records: list, summary: dict, filename: str) -> str:
    os.makedirs(REPORTS_DIR, exist_ok=True)
    path = os.path.join(REPORTS_DIR, filename)
    wb   = Workbook()

    # ── Sheet 1：詳細測試結果 ──
    ws = wb.active
    ws.title = "測試結果"

    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value     = "Slot 遊戲自動化測試報告"
    c.font      = Font(name="Arial", bold=True, size=14, color=_WHITE)
    c.fill      = PatternFill("solid", start_color=_PURPLE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    headers = ["項次","測試時間","玩家資產(前)","押注","遊戲贏分","玩家資產(後)","預期資產","測試結果","失敗原因"]
    for col, h in enumerate(headers, 1):
        _cell(ws, 2, col, h, bg=_PURPLE, bold=True, color=_WHITE)
    ws.row_dimensions[2].height = 22

    for idx, r in enumerate(records):
        row      = idx + 3
        bg       = _GRAY if idx % 2 == 0 else _WHITE
        res_bg   = {"PASS": _GREEN, "FAIL": _RED, "SKIP": _YELLOW}.get(r.pass_fail, _WHITE)
        res_clr  = {"PASS": "375623", "FAIL": "9C0006", "SKIP": "7D4F00"}.get(r.pass_fail, "000000")

        _cell(ws, row, 1, r.index, bg)
        _cell(ws, row, 2, r.test_time, bg)
        _cell(ws, row, 3, r.balance_before, bg)
        _cell(ws, row, 4, r.bet, bg)
        _cell(ws, row, 5, r.win_displayed, bg)
        _cell(ws, row, 6, r.balance_after, bg)
        _cell(ws, row, 7, r.expected_balance if r.expected_balance else "-", bg)
        _cell(ws, row, 8, r.pass_fail, res_bg, bold=True, color=res_clr)
        _cell(ws, row, 9, r.fail_reason, res_bg if r.pass_fail != "PASS" else bg, center=False)
        ws.row_dimensions[row].height = 20

    for i, w in enumerate([6, 22, 14, 8, 10, 14, 12, 10, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Sheet 2：累積金流分析 ──
    ws2 = wb.create_sheet("金流分析")
    ws2.merge_cells("A1:B1")
    c2        = ws2["A1"]
    c2.value  = "累積金流分析"
    c2.font   = Font(name="Arial", bold=True, size=13, color=_WHITE)
    c2.fill   = PatternFill("solid", start_color=_PURPLE)
    c2.alignment = Alignment(horizontal="center")
    ws2.row_dimensions[1].height = 28

    pass_c = sum(1 for r in records if r.pass_fail == "PASS")
    fail_c = sum(1 for r in records if r.pass_fail == "FAIL")
    skip_c = sum(1 for r in records if r.pass_fail == "SKIP")
    stats  = [("PASS 次數", pass_c), ("FAIL 次數", fail_c), ("SKIP 次數", skip_c), *summary.items()]

    for ri, (k, v) in enumerate(stats, 2):
        ck = ws2.cell(ri, 1, k)
        cv = ws2.cell(ri, 2, v)
        ck.font      = Font(name="Arial", bold=True)
        ck.fill      = PatternFill("solid", start_color=_LIGHT_PUR)
        ck.alignment = Alignment(horizontal="left", vertical="center")
        ck.border    = _BD
        is_bad = "FAIL" in str(v) or "差" in str(v)
        cv.fill      = PatternFill("solid", start_color=_RED if is_bad else _GREEN)
        cv.font      = Font(name="Arial", bold=True, color="9C0006" if is_bad else "375623")
        cv.alignment = Alignment(horizontal="center", vertical="center")
        cv.border    = _BD
        ws2.row_dimensions[ri].height = 22

    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 25

    wb.save(path)
    return path
